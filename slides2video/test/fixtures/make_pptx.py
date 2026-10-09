# Builds test/fixtures/mini.pptx (2 slides, speaker notes in the pptx2video-compatible notation, alt-text handles,
# one click animation). Only needed to regenerate the fixture: pip install python-pptx. Apache-2.0, © 2026 天机.
from pptx import Presentation
from pptx.util import Inches, Pt
from lxml import etree
import struct, zlib

def png(w, h, rgb):
    raw = b''.join(b'\x00' + bytes(rgb) * w for _ in range(h))
    def chunk(t, d): return struct.pack('>I', len(d)) + t + d + struct.pack('>I', zlib.crc32(t + d) & 0xffffffff)
    return b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', w, h, 8, 2, 0, 0, 0)) + chunk(b'IDAT', zlib.compress(raw)) + chunk(b'IEND', b'')

open('/tmp/_blue.png', 'wb').write(png(64, 40, (60, 120, 220)))
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
s1 = prs.slides.add_slide(prs.slide_layouts[0])
s1.shapes.title.text = '为什么天空是蓝色的'
s1.placeholders[1].text = '一个公式讲清楚'
s1.notes_slide.notes_text_frame.text = '天空为什么是蓝色的？答案在一个公式里。'
s2 = prs.slides.add_slide(prs.slide_layouts[1])
s2.shapes.title.text = '瑞利散射'
body = s2.placeholders[1].text_frame
body.text = '波长越短，散射越强'
p = body.add_paragraph(); p.text = '蓝光被散射得最多'; p.level = 1
pic = s2.shapes.add_picture('/tmp/_blue.png', Inches(8), Inches(2), Inches(4))
pic._element.nvPicPr.cNvPr.set('descr', '[sky-pic] 蓝天示意')
s2.placeholders[1]._element.nvSpPr.cNvPr.set('descr', '[law]')
s2.notes_slide.notes_text_frame.text = '## [law] 散射定律\n散射强度和波长的四次方成反比。\n## [sky-pic] 看图\n所以天空是蓝色的。[[Spotlight] 看这片蓝天。]'
# one click animation on the picture (Animation Pane order)
spid = pic._element.nvPicPr.cNvPr.get('id')
timing = f'''<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst><p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst><p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst><p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst><p:par><p:cTn id="5" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" nodeType="clickEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst><p:set><p:cBhvr><p:cTn id="6" dur="1" fill="hold"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:seq></p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>'''
s2._element.append(etree.fromstring(timing))
prs.save('mini.pptx')
print('ok mini.pptx')
